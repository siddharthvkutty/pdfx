"""PDF operations. No GUI code here; every function takes and returns plain paths/ints.

Pages in the API are 0-based indexes. Page *specs* (strings like "1,3,5-7") are 1-based,
since that is what users type.
"""
from __future__ import annotations

import glob
import io
import os
import shutil
import subprocess
import sys
import tempfile
import threading
from pathlib import Path

import pypdfium2 as pdfium
from pptx import Presentation
from pptx.util import Emu
from pypdf import PdfReader, PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

PathLike = str | Path

# pdfium is not thread-safe: everything that renders pages (here and in the GUI thumbnails) holds this.
pdfium_lock = threading.Lock()


class PdfxError(Exception):
    """Bad input the user can fix (missing file, encrypted PDF, bad page range)."""


def _open(path: PathLike) -> PdfReader:
    try:
        reader = PdfReader(path)
    except FileNotFoundError:
        raise PdfxError(f"File not found: {path}") from None
    except Exception as e:  # pypdf raises many types for corrupt files
        raise PdfxError(f"Cannot read {path}: {e}") from e
    # shortcut: only PDFs with an empty password are opened, add a password arg when the GUI needs it
    if reader.is_encrypted and not reader.decrypt(""):
        raise PdfxError(f"{path} is password-protected")
    return reader


def _save(writer: PdfWriter, output: PathLike) -> Path:
    # Build in memory first so output may safely be one of the inputs.
    buf = io.BytesIO()
    writer.write(buf)
    out = Path(output)
    out.write_bytes(buf.getvalue())
    return out


def page_count(path: PathLike) -> int:
    return len(_open(path).pages)


def parse_pages(spec: str, page_count: int) -> list[int]:
    """'1,3,5-7' -> [0, 2, 4, 5, 6]. Order and repeats are kept; 'N-M' with N>M counts down."""
    pages: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        try:
            lo, dash, hi = part.partition("-")
            start = int(lo)
            end = int(hi) if dash else start
        except ValueError:
            raise PdfxError(f"Bad page range: {part!r}") from None
        if not (1 <= start <= page_count and 1 <= end <= page_count):
            raise PdfxError(f"Page range {part!r} is outside 1-{page_count}")
        step = 1 if end >= start else -1
        pages.extend(range(start - 1, end - 1 + step, step))
    if not pages:
        raise PdfxError("No pages selected")
    return pages


def merge(inputs: list[PathLike], output: PathLike) -> Path:
    if len(inputs) < 2:
        raise PdfxError("Merge needs at least two PDFs")
    writer = PdfWriter()
    for path in inputs:
        writer.append(_open(path))
    return _save(writer, output)


def reorder(input: PathLike, order: list[int], output: PathLike) -> Path:
    """Write the pages of `input` in `order` (0-based). Omitted pages are dropped, repeats duplicate."""
    reader = _open(input)
    n = len(reader.pages)
    if not order:
        raise PdfxError("No pages selected")
    if any(not 0 <= i < n for i in order):
        raise PdfxError(f"Page index out of range 0-{n - 1}")
    writer = PdfWriter()
    for i in order:
        writer.add_page(reader.pages[i])
    return _save(writer, output)


def extract(input: PathLike, spec: str, output: PathLike) -> Path:
    """Extract pages by 1-based spec, e.g. '4' or '2-5' or '1,3,7-9'."""
    return reorder(input, parse_pages(spec, len(_open(input).pages)), output)


def split(input: PathLike, out_dir: PathLike, specs: list[str] | None = None) -> list[Path]:
    """One output file per spec (default: one per page), named <stem>_<spec>.pdf."""
    reader = _open(input)
    n = len(reader.pages)
    specs = specs or [str(i + 1) for i in range(n)]
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(input).stem
    outputs = []
    for spec in specs:
        writer = PdfWriter()
        for i in parse_pages(spec, n):
            writer.add_page(reader.pages[i])
        outputs.append(_save(writer, out_dir / f"{stem}_{spec.replace(',', '_')}.pdf"))
    return outputs


def _stamp(input: PathLike, output: PathLike, draw) -> Path:
    """Overlay what `draw(canvas, page_w, page_h)` paints on every page of `input`."""
    reader = _open(input)
    writer = PdfWriter()
    for page in reader.pages:
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(w, h))
        draw(c, w, h)
        c.save()
        buf.seek(0)
        # shortcut: pages with /Rotate get the watermark rotated with them, fine for diagonals
        writer.add_page(page).merge_page(PdfReader(buf).pages[0])
    return _save(writer, output)


def watermark_text(
    input: PathLike,
    text: str,
    output: PathLike,
    *,
    opacity: float = 0.3,
    font_size: int = 60,
    angle: float = 45,
    color: tuple[float, float, float] = (0.5, 0.5, 0.5),
) -> Path:
    """Stamp `text` across the centre of every page."""
    if not text:
        raise PdfxError("Watermark text is empty")

    def draw(c, w, h):
        c.setFillColorRGB(*color, alpha=opacity)
        c.setFont("Helvetica", font_size)  # shortcut: built-in font is Latin-1 only, embed a TTF for other scripts
        c.translate(w / 2, h / 2)
        c.rotate(angle)
        c.drawCentredString(0, 0, text)

    return _stamp(input, output, draw)


def watermark_image(
    input: PathLike,
    image: PathLike,
    output: PathLike,
    *,
    opacity: float = 0.3,
    scale: float = 0.5,
    angle: float = 0,
) -> Path:
    """Stamp `image` (PNG/JPG, transparency kept) across the centre; `scale` is its width as a fraction of the page."""
    try:
        img = ImageReader(str(image))
        iw, ih = img.getSize()
    except Exception as e:
        raise PdfxError(f"Cannot read image {image}: {e}") from e

    def draw(c, w, h):
        c.setFillAlpha(opacity)
        dw = w * scale
        dh = dw * ih / iw
        c.translate(w / 2, h / 2)
        c.rotate(angle)
        c.drawImage(img, -dw / 2, -dh / 2, dw, dh, mask="auto")

    return _stamp(input, output, draw)


def _find_tool(names: list[str], windows_globs: list[str], what: str) -> str:
    for n in names:
        if found := shutil.which(n):
            return found
    if sys.platform == "win32":
        for pattern in windows_globs:
            if hits := sorted(glob.glob(pattern)):
                return hits[-1]
    raise PdfxError(f"{what} is not installed or not on PATH")


def _run(cmd: list[str], what: str) -> None:
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    except subprocess.TimeoutExpired:
        raise PdfxError(f"{what} timed out") from None
    if r.returncode != 0:
        raise PdfxError(f"{what} failed: {(r.stderr or r.stdout).strip()}")


def _soffice() -> str:
    return _find_tool(
        ["soffice", "libreoffice"],
        [r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"],
        "LibreOffice",
    )


def libreoffice_path() -> str | None:
    """Where LibreOffice is installed, or None. Cheap: looks only, never starts it."""
    try:
        return _soffice()
    except PdfxError:
        return None


def _soffice_convert(inputs: list[PathLike], target: str, out_dir: Path, infilter: str | None = None) -> None:
    soffice = _soffice()
    out_dir.mkdir(parents=True, exist_ok=True)
    # Private profile: otherwise a running LibreOffice swallows the job or locks the shared profile.
    with tempfile.TemporaryDirectory() as profile:
        _run(
            [soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless",
             *([f"--infilter={infilter}"] if infilter else []),
             "--convert-to", target, "--outdir", str(out_dir), *map(str, inputs)],
            "LibreOffice",
        )


def ppt_to_pdf(inputs: list[PathLike], out_dir: PathLike) -> list[Path]:
    """Convert presentations to <out_dir>/<stem>.pdf with headless LibreOffice."""
    for p in inputs:
        if not Path(p).is_file():
            raise PdfxError(f"File not found: {p}")
    out_dir = Path(out_dir)
    _soffice_convert(inputs, "pdf", out_dir)
    outs = [out_dir / f"{Path(p).stem}.pdf" for p in inputs]
    if missing := [o.name for o in outs if not o.exists()]:
        raise PdfxError(f"LibreOffice produced no PDF for: {', '.join(missing)}")
    return outs


MAX_SLIDE_PT = 4032  # PowerPoint's largest slide is 56 inches


def pdf_to_ppt(input: PathLike, output: PathLike, *, editable: bool = True, dpi: int = 150) -> Path:
    """Convert a PDF to .pptx.

    editable=True: LibreOffice's PDF import, text stays text and shapes stay shapes. Layout can drift on
    complex pages, and scanned pages stay pictures (no OCR).
    editable=False: each page becomes a full-slide picture at `dpi`. Looks identical, nothing is editable.
    """
    _open(input)  # same missing/corrupt/encrypted errors as every other op
    out = Path(output)
    if editable:
        with tempfile.TemporaryDirectory() as tmp:
            _soffice_convert([input], "pptx", Path(tmp), infilter="impress_pdf_import")
            made = Path(tmp) / f"{Path(input).stem}.pptx"
            if not made.exists():
                raise PdfxError("LibreOffice produced no PowerPoint file")
            shutil.move(made, out)
        return out

    prs = Presentation()
    with pdfium_lock:
        pdf = pdfium.PdfDocument(str(input))
    try:
        sizes = [pdf[i].get_size() for i in range(len(pdf))] if len(pdf) else []
        if not sizes:
            raise PdfxError("PDF has no pages")
        w0, h0 = sizes[0]
        if max(w0, h0) > MAX_SLIDE_PT:
            raise PdfxError("Pages are larger than PowerPoint's 56-inch slide limit")
        pt = lambda v: Emu(int(v * 12700))
        prs.slide_width, prs.slide_height = pt(w0), pt(h0)  # shortcut: slide size follows page 1, other sizes are letterboxed
        for i, (pw, ph) in enumerate(sizes):
            with pdfium_lock:
                img = pdf[i].render(scale=dpi / 72).to_pil()
            buf = io.BytesIO()
            img.save(buf, "PNG")  # shortcut: PNG is large for photo pages, offer JPEG if decks get too big
            fit = min(w0 / pw, h0 / ph)
            w, h = pw * fit, ph * fit
            prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(
                buf, pt((w0 - w) / 2), pt((h0 - h) / 2), pt(w), pt(h)
            )
    finally:
        with pdfium_lock:
            pdf.close()
    prs.save(out)
    return out


COMPRESSION_LEVELS = ("printer", "ebook", "screen")  # least -> most compression (Ghostscript -dPDFSETTINGS)


def compress(input: PathLike, output: PathLike, level: str = "ebook") -> Path:
    """Recompress with Ghostscript. If that does not make the file smaller, the original bytes are written instead."""
    if level not in COMPRESSION_LEVELS:
        raise PdfxError(f"Unknown compression level {level!r}")
    _open(input)  # same missing/corrupt/encrypted errors as every other op
    gs = _find_tool(
        ["gs", "gswin64c", "gswin32c"],
        [r"C:\Program Files\gs\gs*\bin\gswin64c.exe", r"C:\Program Files (x86)\gs\gs*\bin\gswin32c.exe"],
        "Ghostscript",
    )
    out = Path(output)
    fd, tmp = tempfile.mkstemp(suffix=".pdf", dir=out.parent)  # same dir so os.replace is atomic
    os.close(fd)
    try:
        _run(
            [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4", f"-dPDFSETTINGS=/{level}",
             "-dNOPAUSE", "-dQUIET", "-dBATCH", f"-sOutputFile={tmp}", str(input)],
            "Ghostscript",
        )
        if Path(tmp).stat().st_size >= Path(input).stat().st_size:
            shutil.copyfile(input, tmp)
        os.replace(tmp, out)
    finally:
        Path(tmp).unlink(missing_ok=True)
    return out
