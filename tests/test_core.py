import shutil

import pytest
from pypdf import PdfReader
from reportlab.pdfgen import canvas

import pdfx
from pdfx import PdfxError


def make_pdf(path, label, pages):
    c = canvas.Canvas(str(path))
    for i in range(1, pages + 1):
        c.drawString(100, 700, f"{label}{i}")
        c.showPage()
    c.save()
    return path


def labels(path):
    return [p.extract_text().strip() for p in PdfReader(path).pages]


@pytest.fixture
def a(tmp_path):
    return make_pdf(tmp_path / "a.pdf", "A", 5)


@pytest.fixture
def b(tmp_path):
    return make_pdf(tmp_path / "b.pdf", "B", 2)


def test_merge(a, b, tmp_path):
    out = pdfx.merge([a, b], tmp_path / "m.pdf")
    assert labels(out) == ["A1", "A2", "A3", "A4", "A5", "B1", "B2"]


def test_merge_needs_two(a, tmp_path):
    with pytest.raises(PdfxError):
        pdfx.merge([a], tmp_path / "m.pdf")


def test_reorder_and_drop(a, tmp_path):
    out = pdfx.reorder(a, [4, 0, 2], tmp_path / "r.pdf")
    assert labels(out) == ["A5", "A1", "A3"]


def test_reorder_out_of_range(a, tmp_path):
    with pytest.raises(PdfxError):
        pdfx.reorder(a, [5], tmp_path / "r.pdf")


def test_in_place_overwrite(a):
    pdfx.reorder(a, [4, 3], a)
    assert labels(a) == ["A5", "A4"]


@pytest.mark.parametrize(
    "spec,expected",
    [("3", [2]), ("2-4", [1, 2, 3]), ("1,3-4,5", [0, 2, 3, 4]), ("4-2", [3, 2, 1]), (" 1 , 2 ", [0, 1])],
)
def test_parse_pages(spec, expected):
    assert pdfx.parse_pages(spec, 5) == expected


@pytest.mark.parametrize("spec", ["0", "6", "1-9", "x", "1-", "", ","])
def test_parse_pages_bad(spec):
    with pytest.raises(PdfxError):
        pdfx.parse_pages(spec, 5)


def test_extract_single_and_range(a, tmp_path):
    assert labels(pdfx.extract(a, "2", tmp_path / "s.pdf")) == ["A2"]
    assert labels(pdfx.extract(a, "2-4", tmp_path / "r.pdf")) == ["A2", "A3", "A4"]


def test_split_every_page(a, tmp_path):
    outs = pdfx.split(a, tmp_path / "out")
    assert [labels(o) for o in outs] == [[f"A{i}"] for i in range(1, 6)]


def test_split_ranges(a, tmp_path):
    outs = pdfx.split(a, tmp_path / "out", ["1-2", "3,5"])
    assert [o.name for o in outs] == ["a_1-2.pdf", "a_3_5.pdf"]
    assert [labels(o) for o in outs] == [["A1", "A2"], ["A3", "A5"]]


def test_watermark(a, tmp_path):
    out = pdfx.watermark_text(a, "DRAFT", tmp_path / "w.pdf")
    pages = PdfReader(out).pages
    assert len(pages) == 5
    assert all("DRAFT" in p.extract_text() and f"A{i}" in p.extract_text() for i, p in enumerate(pages, 1))


def test_watermark_empty(a, tmp_path):
    with pytest.raises(PdfxError):
        pdfx.watermark_text(a, "", tmp_path / "w.pdf")


def test_missing_and_corrupt(tmp_path):
    with pytest.raises(PdfxError):
        pdfx.extract(tmp_path / "nope.pdf", "1", tmp_path / "o.pdf")
    bad = tmp_path / "bad.pdf"
    bad.write_text("not a pdf")
    with pytest.raises(PdfxError):
        pdfx.extract(bad, "1", tmp_path / "o.pdf")


def test_encrypted(a, tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter(clone_from=a)
    w.encrypt("secret")
    enc = tmp_path / "enc.pdf"
    w.write(enc)
    with pytest.raises(PdfxError, match="password"):
        pdfx.extract(enc, "1", tmp_path / "o.pdf")


def test_page_count(a):
    assert pdfx.page_count(a) == 5


def test_watermark_image(a, tmp_path):
    from PIL import Image

    png = tmp_path / "logo.png"
    Image.new("RGBA", (40, 20), (255, 0, 0, 128)).save(png)
    out = pdfx.watermark_image(a, png, tmp_path / "w.pdf", opacity=0.5)
    pages = PdfReader(out).pages
    assert len(pages) == 5
    assert all(len(p.images) == 1 and f"A{i}" in p.extract_text() for i, p in enumerate(pages, 1))


def test_watermark_image_bad(a, tmp_path):
    bad = tmp_path / "x.png"
    bad.write_text("nope")
    with pytest.raises(PdfxError):
        pdfx.watermark_image(a, bad, tmp_path / "w.pdf")


needs_gs = pytest.mark.skipif(not shutil.which("gs"), reason="Ghostscript not installed")
needs_office = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


@needs_gs
def test_compress(tmp_path):
    from PIL import Image

    img = tmp_path / "big.png"
    Image.effect_noise((1500, 1500), 60).convert("RGB").save(img)
    src = tmp_path / "big.pdf"
    c = canvas.Canvas(str(src))
    c.drawImage(str(img), 0, 0, 500, 500)
    c.save()
    out = pdfx.compress(src, tmp_path / "small.pdf", "screen")
    assert out.stat().st_size < src.stat().st_size
    assert len(PdfReader(out).pages) == 1


@needs_gs
def test_compress_never_grows(a, tmp_path):
    out = pdfx.compress(a, tmp_path / "c.pdf")
    assert out.stat().st_size <= a.stat().st_size
    assert labels(out) == [f"A{i}" for i in range(1, 6)]


def test_compress_bad_level(a, tmp_path):
    with pytest.raises(PdfxError):
        pdfx.compress(a, tmp_path / "c.pdf", "max")


@needs_office
def test_ppt_to_pdf(tmp_path):
    odp = tmp_path / "deck.fodp"  # flat ODF presentation: plain XML, no binary fixture needed
    odp.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'office:version="1.2" office:mimetype="application/vnd.oasis.opendocument.presentation">'
        "<office:body><office:presentation><draw:page draw:name='one'><draw:frame>"
        "<draw:text-box><text:p>HELLOSLIDE</text:p></draw:text-box></draw:frame></draw:page>"
        "</office:presentation></office:body></office:document>"
    )
    [out] = pdfx.ppt_to_pdf([odp], tmp_path / "out")
    assert out.name == "deck.pdf"
    assert "HELLOSLIDE" in PdfReader(out).pages[0].extract_text()


def test_ppt_missing_file(tmp_path):
    with pytest.raises(PdfxError):
        pdfx.ppt_to_pdf([tmp_path / "nope.pptx"], tmp_path)


def test_pdf_to_ppt_images(tmp_path):
    from pptx import Presentation

    src = tmp_path / "wide.pdf"
    c = canvas.Canvas(str(src), pagesize=(960, 540))
    c.drawString(100, 400, "SLIDE ONE")
    c.showPage()
    c.setPageSize((540, 960))  # a differently shaped second page gets letterboxed, not stretched
    c.drawString(100, 400, "SLIDE TWO")
    c.showPage()
    c.save()
    out = pdfx.pdf_to_ppt(src, tmp_path / "o.pptx", editable=False, dpi=72)
    prs = Presentation(out)
    assert len(prs.slides) == 2
    assert (prs.slide_width, prs.slide_height) == (960 * 12700, 540 * 12700)
    for slide in prs.slides:
        [pic] = slide.shapes
        assert pic.shape_type == 13  # PICTURE
        assert pic.top >= 0 and pic.left >= 0 and pic.height <= prs.slide_height and pic.width <= prs.slide_width
    tall = prs.slides[1].shapes[0]
    assert abs(tall.width / tall.height - 540 / 960) < 0.01  # aspect kept


def test_pdf_to_ppt_too_big(tmp_path):
    src = tmp_path / "huge.pdf"
    c = canvas.Canvas(str(src), pagesize=(5000, 3000))
    c.showPage()
    c.save()
    with pytest.raises(PdfxError, match="56-inch"):
        pdfx.pdf_to_ppt(src, tmp_path / "o.pptx", editable=False)


@needs_office
def test_pdf_to_ppt_editable_keeps_text(tmp_path):
    from pptx import Presentation

    src = tmp_path / "s.pdf"
    c = canvas.Canvas(str(src), pagesize=(960, 540))
    c.setFont("Helvetica", 30)
    c.drawString(100, 400, "Editable headline")
    c.showPage()
    c.save()
    out = pdfx.pdf_to_ppt(src, tmp_path / "o.pptx")
    texts = [sh.text_frame.text for sh in Presentation(out).slides[0].shapes if sh.has_text_frame]
    assert "Editable headline" in texts


def test_pdf_to_ppt_bad_input(tmp_path):
    with pytest.raises(PdfxError):
        pdfx.pdf_to_ppt(tmp_path / "nope.pdf", tmp_path / "o.pptx")


def _image(path, size=(60, 30), color=(200, 30, 30), **save):
    from PIL import Image

    Image.new("RGB", size, color).save(path, **save)
    return path


@pytest.mark.parametrize("name", ["a.jpg", "a.jpeg", "a.png", "a.gif", "a.bmp", "a.tif", "a.tiff", "a.webp"])
def test_images_to_pdf_common_formats(tmp_path, name):
    out = pdfx.images_to_pdf([_image(tmp_path / name)], tmp_path / "o.pdf")
    assert len(PdfReader(out).pages) == 1


def test_images_to_pdf_order_and_page_size_from_dpi(tmp_path):
    first = _image(tmp_path / "a.png", (600, 300), dpi=(300, 300))  # 2 x 1 inch
    second = _image(tmp_path / "b.png", (600, 300))  # no DPI: 150 assumed, 4 x 2 inches
    out = pdfx.images_to_pdf([second, first], tmp_path / "o.pdf")
    sizes = [(round(p.mediabox.width), round(p.mediabox.height)) for p in PdfReader(out).pages]  # PNG stores DPI as pixels per metre, so 300 reads back as 299.9994
    assert sizes == [(288, 144), (144, 72)]


def test_images_to_pdf_keeps_jpeg_bytes_compressed(tmp_path):
    from PIL import Image

    photo = tmp_path / "p.jpg"
    Image.effect_noise((1200, 900), 80).convert("RGB").save(photo, quality=85)
    out = pdfx.images_to_pdf([photo], tmp_path / "o.pdf")
    assert PdfReader(out).pages[0].images[0].name.endswith(".jpg")
    assert out.stat().st_size < photo.stat().st_size * 1.2  # embedded as-is, not re-encoded larger


@pytest.mark.parametrize("orientation,rotate", [(1, 0), (3, 180), (6, 90), (8, 270)])
def test_images_to_pdf_exif_rotation(tmp_path, orientation, rotate):
    import pypdfium2 as pdfium
    from PIL import Image

    img = Image.new("RGB", (60, 30), (0, 0, 255))
    exif = img.getexif()
    exif[0x0112] = orientation
    img.save(tmp_path / "p.jpg", exif=exif, dpi=(150, 150))
    out = pdfx.images_to_pdf([tmp_path / "p.jpg"], tmp_path / "o.pdf")
    page = PdfReader(out).pages[0]
    assert (page.get("/Rotate") or 0) == rotate
    assert (round(page.mediabox.width), round(page.mediabox.height)) == (29, 14)  # the box matches the stored pixels
    shown = pdfium.PdfDocument(str(out))[0].render(scale=4).to_pil()
    # displayed upright: 90/270 swap width and height; the image must fill the whole page, no cropping or blank strips
    assert (shown.width > shown.height) == (rotate in (0, 180))
    for corner in ((2, 2), (shown.width - 3, 2), (2, shown.height - 3), (shown.width - 3, shown.height - 3)):
        assert shown.getpixel(corner)[2] > 200 and shown.getpixel(corner)[0] < 60


def test_images_to_pdf_transparency_is_white_not_black(tmp_path):
    import pypdfium2 as pdfium
    from PIL import Image

    Image.new("RGBA", (40, 40), (0, 0, 0, 0)).save(tmp_path / "clear.png")  # fully transparent
    out = pdfx.images_to_pdf([tmp_path / "clear.png"], tmp_path / "o.pdf")
    pixel = pdfium.PdfDocument(str(out))[0].render(scale=1).to_pil().getpixel((5, 5))
    assert pixel[:3] == (255, 255, 255)


def test_images_to_pdf_errors(tmp_path):
    with pytest.raises(PdfxError, match="No images"):
        pdfx.images_to_pdf([], tmp_path / "o.pdf")
    bad = tmp_path / "x.png"
    bad.write_text("not an image")
    with pytest.raises(PdfxError, match="Cannot read image"):
        pdfx.images_to_pdf([bad], tmp_path / "o.pdf")
    with pytest.raises(PdfxError, match="Cannot read image"):
        pdfx.images_to_pdf([tmp_path / "missing.jpg"], tmp_path / "o.pdf")
